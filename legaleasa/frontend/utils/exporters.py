
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document


def export_txt(text: str, title: str = "LegalEase Document") -> bytes:
    return text.encode("utf-8")


def export_pdf(text: str, title: str = "LegalEase Document") -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    x_margin = 50
    y = height - 70
    line_height = 16

    pdf.setTitle(title)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(x_margin, y, title[:70])
    y -= 35
    pdf.setFont("Helvetica", 10)

    for paragraph in text.splitlines():
        words = paragraph.split()
        line = ""

        for word in words:
            test_line = f"{line} {word}".strip()

            if pdf.stringWidth(test_line, "Helvetica", 10) > width - 2 * x_margin:
                if line:
                    pdf.drawString(x_margin, y, line)
                    y -= line_height
                line = word
            else:
                line = test_line

            if y < 50:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 50

        if line:
            pdf.drawString(x_margin, y, line)
            y -= line_height + 5

        if y < 50:
            pdf.showPage()
            pdf.setFont("Helvetica", 10)
            y = height - 50

    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()


def export_docx(text: str, title: str = "LegalEase Document") -> bytes:
    document = Document()
    document.add_heading(title, level=1)

    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()