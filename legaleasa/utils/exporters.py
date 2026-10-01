
from io import BytesIO
from docx import Document
from reportlab.pdfgen import canvas


def export_txt(text):
    return text.encode("utf-8")


def export_docx(text):
    doc = Document()
    doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_pdf(text):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)

    y = 800
    for line in text.split("\n"):
        if y < 50:
            pdf.showPage()
            y = 800
        pdf.drawString(50, y, line[:100])
        y -= 20

    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()
    from io import BytesIO


def export_txt(text: str) -> bytes:
    """Convert text into a downloadable TXT file."""
    return text.encode("utf-8")


def export_pdf(text: str) -> bytes:
    """Convert text into a downloadable PDF file."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)

    width, height = A4
    x_margin = 50
    y = height - 50
    line_height = 16

    pdf.setFont("Helvetica", 10)

    for paragraph in text.splitlines():
        words = paragraph.split()
        line = ""

        for word in words:
            test_line = f"{line} {word}".strip()

            if pdf.stringWidth(test_line, "Helvetica", 10) > width - 2 * x_margin:
                pdf.drawString(x_margin, y, line)
                y -= line_height
                line = word
            else:
                line = test_line

            if y < 50:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 50

        if line:
            pdf.drawString(x_margin, y, line)
            y -= line_height

        y -= 5

        if y < 50:
            pdf.showPage()
            pdf.setFont("Helvetica", 10)
            y = height - 50

    pdf.save()
    buffer.seek(0)
    return buffer.getvalue()


def export_docx(text: str) -> bytes:
    """Convert text into a downloadable DOCX file."""
    from docx import Document

    document = Document()

    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()